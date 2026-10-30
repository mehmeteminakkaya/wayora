from __future__ import annotations

import html
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "professor"
OUT = SRC

INK = RGBColor(0x12, 0x12, 0x12)
INK_2 = RGBColor(0x55, 0x57, 0x5C)
LINE = "DEDED9"
SIGNAL = "FFC400"
BG = "F3F3F1"
SURFACE = "FFFFFF"

BODY_FONT = "Arial"
DISPLAY_FONT = "Arial Black"
MONO_FONT = "Consolas"
CHROME = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")

DOCS = [
    ("00-Sunum-Rehberi.md", "00-Sunum-Rehberi", "Wayora-Sunum-Rehberi"),
    ("01-Yapilanlar.md", "01-Yapilanlar", "Wayora-Yapilanlar"),
    ("02-User-Stories.md", "02-User-Stories", "Wayora-User-Stories"),
    ("03-Proje-Yonetimi.md", "03-Proje-Yonetimi", "Wayora-Proje-Yonetimi"),
    ("04-Ders-Konseptleri.md", "04-Ders-Konseptleri", "Wayora-Ders-Konseptleri"),
    ("05-Architecture.md", "05-Architecture", "Wayora-Architecture"),
    ("06-UI-UX.md", "06-UI-UX", "Wayora-UI-UX"),
]


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in [("w:top", top), ("w:start", start), ("w:bottom", bottom), ("w:end", end)]:
        node = OxmlElement(tag)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def frontmatter(text: str) -> tuple[dict[str, str], str]:
    if not text.startswith("---"):
        return {}, text
    parts = text.split("---", 2)
    if len(parts) < 3:
        return {}, text
    meta: dict[str, str] = {}
    for line in parts[1].strip().splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            meta[key.strip()] = value.strip()
    return meta, parts[2].lstrip()


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, font, color, before, after in [
        ("Heading 1", 18, DISPLAY_FONT, INK, 14, 8),
        ("Heading 2", 13, BODY_FONT, INK, 13, 5),
        ("Heading 3", 11, BODY_FONT, INK_2, 10, 3),
    ]:
        style = styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "WAYORA  /  JOB-191"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = DISPLAY_FONT
        run.font.size = Pt(7.5)
        run.font.color.rgb = INK_2
    header.add_run("\t")
    tail = header.add_run("PROFESÖR TESLİMİ · 2026")
    tail.font.name = MONO_FONT
    tail.font.size = Pt(6.8)
    tail.font.color.rgb = INK_2

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Wayora · İstinye Üniversitesi · ")
    run.font.name = MONO_FONT
    run.font.size = Pt(7)
    run.font.color.rgb = INK_2
    add_page_field(footer)
    return doc


def add_cover(doc: Document, meta: dict[str, str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("WAYORA")
    r.font.name = DISPLAY_FONT
    r.font.size = Pt(34)
    r.font.color.rgb = INK
    r2 = p.add_run(" / ")
    r2.font.name = DISPLAY_FONT
    r2.font.size = Pt(34)
    r2.font.color.rgb = INK_2
    r3 = p.add_run("JOB-191")
    r3.font.name = MONO_FONT
    r3.font.size = Pt(11)
    r3.font.color.rgb = INK_2

    rule = doc.add_table(rows=1, cols=1)
    rule.autofit = False
    rule.columns[0].width = Cm(16.8)
    rule.cell(0, 0).width = Cm(16.8)
    shade(rule.cell(0, 0), SIGNAL)
    rule.cell(0, 0).text = ""
    rule.cell(0, 0).paragraphs[0].paragraph_format.space_after = Pt(0)
    marker = rule.cell(0, 0).paragraphs[0].add_run(" ")
    marker.font.size = Pt(3)
    marker.font.color.rgb = RGBColor(0xFF, 0xC4, 0x00)
    doc.add_paragraph()

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(22)
    r = title.add_run(meta.get("title", "Wayora Belgesi").upper())
    r.font.name = DISPLAY_FONT
    r.font.size = Pt(26)
    r.font.color.rgb = INK

    subtitle = doc.add_paragraph()
    r = subtitle.add_run(meta.get("subtitle", ""))
    r.font.name = BODY_FONT
    r.font.size = Pt(13)
    r.font.color.rgb = INK_2

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Table Grid"
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    rows = [
        ("Belge türü", meta.get("doc_type", "TESLİM")),
        ("Sürüm", meta.get("version", "1.0")),
        ("Tarih", meta.get("date", "30 Ekim 2026")),
        ("Ders", "JOB-191 · AI-Augmented Software Engineering"),
    ]
    for index, (key, value) in enumerate(rows):
        meta_table.cell(index, 0).text = key
        meta_table.cell(index, 1).text = value
        shade(meta_table.cell(index, 0), BG)
        for cell in meta_table.rows[index].cells:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = BODY_FONT
                    run.font.size = Pt(9)
                    run.font.color.rgb = INK
        meta_table.cell(index, 0).paragraphs[0].runs[0].font.bold = True
        meta_table.cell(index, 0).width = Cm(4.5)
        meta_table.cell(index, 1).width = Cm(12.3)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(20)
    r = note.add_run("Wayora ekibi tarafından hazırlanmıştır. Canlı ürün: wayora-ten.vercel.app")
    r.font.name = MONO_FONT
    r.font.size = Pt(8)
    r.font.color.rgb = INK_2
    doc.add_page_break()


def add_inline(paragraph, text: str, size=10.3, color=INK) -> None:
    # Simple inline formatting: **bold**, `code`.
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(max(8.5, size - 0.8))
            run.font.color.rgb = INK_2
        else:
            run = paragraph.add_run(part)
        run.font.name = BODY_FONT
        if not (part.startswith("`") and part.endswith("`")):
            run.font.size = Pt(size)
            run.font.color.rgb = color


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        add_inline(paragraph, header, size=8.4, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "121212")
        set_cell_margins(cell)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row[: len(headers)]):
            cells[index].text = ""
            add_inline(cells[index].paragraphs[0], value, size=8.2)
            set_cell_margins(cells[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(source: Path, output: Path) -> None:
    meta, body = frontmatter(source.read_text(encoding="utf-8"))
    doc = setup_document()
    add_cover(doc, meta)
    lines = body.splitlines()
    index = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if not stripped or stripped == "---":
            index += 1
            continue
        image_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if image_match:
            caption, relative_path = image_match.groups()
            image_path = (source.parent / relative_path).resolve()
            if image_path.exists():
                if "-mobile" in relative_path:
                    width = Cm(7.2)
                elif "architecture" in relative_path:
                    width = Cm(16.2)
                else:
                    width = Cm(15.0)
                doc.add_picture(str(image_path), width=width)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if caption:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = cap.add_run(caption)
                    run.font.name = MONO_FONT
                    run.font.size = Pt(7.6)
                    run.font.color.rgb = INK_2
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\s*\|?[\s:|-]+\|", lines[index + 1]
        ):
            headers = [cell.strip() for cell in stripped.strip("|").split("|")]
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(doc, headers, rows)
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            doc.add_heading(text, level=min(level, 3))
            index += 1
            continue
        if stripped.startswith("- "):
            items = []
            while index < len(lines) and lines[index].strip().startswith("- "):
                items.append(lines[index].strip()[2:])
                index += 1
            for item in items:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, item, size=10)
            continue
        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while index < len(lines) and re.match(r"^\d+\.\s+", lines[index].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[index].strip()))
                index += 1
            for item in items:
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, item, size=10)
            continue
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            table = doc.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            cell = table.cell(0, 0)
            cell.text = ""
            add_inline(cell.paragraphs[0], text, size=9.5)
            shade(cell, BG)
            set_cell_margins(cell, 100, 120, 100, 120)
            doc.add_paragraph()
            index += 1
            continue
        # paragraph: collect until blank or structural line
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate == "---" or candidate.startswith(("#", "|", "- ", ">")):
                break
            paragraph_lines.append(candidate)
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))
    doc.save(output)
    print("wrote", output)


def inline_html(text: str) -> str:
    escaped = html.escape(text)
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", escaped)
    escaped = re.sub(r"`([^`]+)`", r"<code>\1</code>", escaped)
    return escaped


def build_html(source: Path, output: Path) -> None:
    meta, body = frontmatter(source.read_text(encoding="utf-8"))
    lines = body.splitlines()
    parts: list[str] = []
    index = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if not stripped or stripped == "---":
            index += 1
            continue
        image_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if image_match:
            caption, relative_path = image_match.groups()
            parts.append(
                f'<figure><img src="{html.escape(relative_path)}" alt="{html.escape(caption)}">'
                f'<figcaption>{inline_html(caption)}</figcaption></figure>'
            )
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\s*\|?[\s:|-]+\|", lines[index + 1]
        ):
            headers = [cell.strip() for cell in stripped.strip("|").split("|")]
            parts.append("<table><thead><tr>")
            for header in headers:
                parts.append(f"<th>{inline_html(header)}</th>")
            parts.append("</tr></thead><tbody>")
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                parts.append("<tr>")
                for cell in row:
                    parts.append(f"<td>{inline_html(cell)}</td>")
                parts.append("</tr>")
                index += 1
            parts.append("</tbody></table>")
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            parts.append(f"<h{min(level, 3)}>{inline_html(text)}</h{min(level, 3)}>")
            index += 1
            continue
        if stripped.startswith("- "):
            parts.append("<ul>")
            while index < len(lines) and lines[index].strip().startswith("- "):
                parts.append(f"<li>{inline_html(lines[index].strip()[2:])}</li>")
                index += 1
            parts.append("</ul>")
            continue
        if re.match(r"^\d+\.\s+", stripped):
            parts.append("<ol>")
            while index < len(lines) and re.match(r"^\d+\.\s+", lines[index].strip()):
                item = re.sub(r"^\d+\.\s+", "", lines[index].strip())
                parts.append(f"<li>{inline_html(item)}</li>")
                index += 1
            parts.append("</ol>")
            continue
        if stripped.startswith(">"):
            parts.append(f"<blockquote>{inline_html(stripped.lstrip('> ').strip())}</blockquote>")
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate == "---" or candidate.startswith(("#", "|", "- ", ">")):
                break
            paragraph_lines.append(candidate)
            index += 1
        parts.append(f"<p>{inline_html(' '.join(paragraph_lines))}</p>")

    title = html.escape(meta.get("title", "Wayora Belgesi"))
    subtitle = html.escape(meta.get("subtitle", ""))
    doc_type = html.escape(meta.get("doc_type", "TESLİM"))
    version = html.escape(meta.get("version", "1.0"))
    date = html.escape(meta.get("date", "30 Ekim 2026"))
    page = f"""<!doctype html>
<html lang="tr">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  @page {{ size: A4; margin: 18mm 17mm 17mm; }}
  * {{ box-sizing: border-box; }}
  body {{ margin: 0; color: #121212; background: #fff; font-family: Arial, Helvetica, sans-serif; font-size: 10.5pt; line-height: 1.5; }}
  h1 {{ font-family: "Arial Black", Arial, sans-serif; font-size: 24pt; line-height: 1.08; margin: 0 0 8pt; }}
  h2 {{ font-size: 14pt; margin: 18pt 0 6pt; page-break-after: avoid; break-after: avoid; }}
  h3 {{ font-size: 11.5pt; margin: 12pt 0 4pt; color: #444; page-break-after: avoid; break-after: avoid; }}
  p {{ margin: 0 0 7pt; }}
  ul, ol {{ margin: 0 0 8pt 18pt; padding: 0; }}
  li {{ margin: 0 0 3pt; }}
  strong {{ font-weight: 700; }}
  code {{ font-family: Consolas, monospace; font-size: 9pt; background: #f3f3f1; padding: 1px 4px; }}
  blockquote {{ margin: 10pt 0; padding: 9pt 11pt; border-left: 4px solid #ffc400; background: #f3f3f1; }}
  table {{ width: 100%; border-collapse: collapse; margin: 10pt 0 14pt; font-size: 8.8pt; page-break-inside: auto; }}
  tr {{ page-break-inside: avoid; }}
  th {{ background: #121212; color: #fff; text-align: left; padding: 6pt 7pt; border: 1px solid #121212; }}
  td {{ padding: 5pt 7pt; border: 1px solid #deded9; vertical-align: top; }}
  tbody tr:nth-child(even) {{ background: #fafafa; }}
  figure {{ margin: 14pt 0 16pt; page-break-inside: avoid; }}
  figure img {{ display: block; width: auto; max-width: 100%; max-height: 205mm; margin: 0 auto; border: 1px solid #deded9; }}
  figcaption {{ margin-top: 5pt; text-align: center; color: #55575c; font-family: Consolas, monospace; font-size: 7.8pt; }}
  .cover {{ min-height: 245mm; page-break-after: always; padding-top: 22mm; }}
  .brand {{ font-family: "Arial Black", Arial, sans-serif; font-size: 26pt; letter-spacing: -0.02em; }}
  .brand span {{ color: #777; font-family: Consolas, monospace; font-size: 9pt; letter-spacing: 0; }}
  .rule {{ width: 100%; height: 8px; background: #ffc400; margin: 16pt 0 26pt; }}
  .cover h1 {{ font-size: 34pt; max-width: 155mm; }}
  .subtitle {{ font-size: 14pt; color: #55575c; max-width: 150mm; }}
  .meta {{ margin-top: 28mm; border-top: 1px solid #deded9; padding-top: 8pt; color: #55575c; font-family: Consolas, monospace; font-size: 8.5pt; }}
  .meta b {{ color: #121212; }}
  .page-break {{ page-break-before: always; }}
</style>
</head>
<body>
<section class="cover">
  <div class="brand">WAYORA <span>/ JOB-191</span></div>
  <div class="rule"></div>
  <h1>{title}</h1>
  <p class="subtitle">{subtitle}</p>
  <div class="meta"><b>{doc_type}</b> · Sürüm {version} · {date}<br>
  JOB-191 · AI-Augmented Software Engineering<br>
  Canlı ürün: wayora-ten.vercel.app</div>
</section>
{''.join(parts)}
</body>
</html>"""
    output.write_text(page, encoding="utf-8")


def build_pdf(html_path: Path, pdf_path: Path) -> bool:
    if not CHROME.exists():
        return False
    profile = Path(tempfile.mkdtemp(prefix="wayora-chrome-"))
    try:
        result = subprocess.run(
            [
                str(CHROME),
                "--headless=new",
                "--disable-gpu",
                "--no-sandbox",
                "--no-pdf-header-footer",
                f"--user-data-dir={profile}",
                f"--print-to-pdf={pdf_path}",
                html_path.as_uri(),
            ],
            capture_output=True,
            text=True,
            timeout=120,
        )
        if result.returncode != 0:
            print("pdf warning", pdf_path.name, result.stderr[-300:])
        return pdf_path.exists()
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def main() -> None:
    for source_name, stem, root_name in DOCS:
        source = SRC / source_name
        if not source.exists():
            print("missing", source)
            continue
        docx_path = OUT / f"{stem}.docx"
        html_path = OUT / f"{stem}.html"
        pdf_path = OUT / f"{stem}.pdf"
        build_docx(source, docx_path)
        build_html(source, html_path)
        build_pdf(html_path, pdf_path)
        html_path.unlink(missing_ok=True)
        # Eski kök dosya adlarını da güncel tut.
        if root_name in {"Wayora-User-Stories", "Wayora-Proje-Yonetimi"}:
            try:
                shutil.copyfile(docx_path, ROOT / "docs" / f"{root_name}.docx")
                if pdf_path.exists():
                    shutil.copyfile(pdf_path, ROOT / "docs" / f"{root_name}.pdf")
            except OSError as error:
                print("warning: root copy skipped", root_name, error)
        print("ready", docx_path, "pdf", pdf_path.exists())


if __name__ == "__main__":
    main()
