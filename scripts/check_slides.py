from pptx import Presentation

prs = Presentation('docs/Wayora-Sunum.pptx')
width = prs.slide_width
height = prs.slide_height
print('slides:', len(prs.slides), 'size:', width, height)
for index, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        try:
            right = shape.left + shape.width
            bottom = shape.top + shape.height
        except Exception:
            continue
        if shape.left < 0 or shape.top < 0 or right > width + 5000 or bottom > height + 5000:
            print(f'slide {index}: outside -> {shape.shape_type} {shape.name} left={shape.left} top={shape.top} right={right} bottom={bottom}')
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
    print(f'slide {index}: shapes={len(slide.shapes)} notes={len(notes)}')
